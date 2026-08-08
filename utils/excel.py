"""Foydalanuvchilar ro'yxatini Excel'ga yozish.

Fayl doim bitta: exports/users.xlsx
Bor bo'lsa -> yangilanadi, yo'q bo'lsa -> yaratiladi. Har kuni yangi fayl ochilmaydi.
"""

import asyncio
import logging
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import EXPORT_FILE
from db.queries import TZ, get_all_users

logger = logging.getLogger(__name__)

HEADERS = ["№", "Telegram ID", "Ism", "Username", "Telefon", "Holat", "Ro'yxatdan o'tgan"]
WIDTHS = [6, 16, 28, 20, 18, 12, 22]


def _write(rows: list[list], path: Path) -> None:
    if path.exists():
        try:
            wb = load_workbook(path)
            ws = wb.active
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)  # sarlavhani qoldirib tozalaymiz
        except Exception:  # fayl buzilgan bo'lsa qaytadan yaratamiz
            logger.warning("Excel fayl o'qilmadi, qaytadan yaratiladi")
            wb = Workbook()
            ws = wb.active
            ws.title = "Foydalanuvchilar"
            ws.append(HEADERS)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Foydalanuvchilar"
        ws.append(HEADERS)

    if ws.max_row == 0:
        ws.append(HEADERS)

    header_fill = PatternFill("solid", start_color="4F81BD")
    for col, width in enumerate(WIDTHS, start=1):
        cell = ws.cell(row=1, column=col)
        cell.value = HEADERS[col - 1]
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    for row in rows:
        ws.append(row)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


async def build_excel() -> Path:
    users = await get_all_users()
    rows = []
    for i, u in enumerate(users, start=1):
        created = u.created_at.astimezone(TZ).strftime("%Y-%m-%d %H:%M") if u.created_at else ""
        rows.append(
            [
                i,
                u.tg_id,
                u.full_name or "",
                f"@{u.username}" if u.username else "",
                u.phone or "",
                "Faol" if u.is_active else "Bloklagan",
                created,
            ]
        )
    await asyncio.to_thread(_write, rows, EXPORT_FILE)
    logger.info("Excel yangilandi: %s (%s ta foydalanuvchi)", EXPORT_FILE, len(rows))
    return EXPORT_FILE
